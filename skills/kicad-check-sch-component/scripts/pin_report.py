#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pin_report.py - 符号校验的引脚比对表（xlsx）。

    pin | 手册名 | 网表名 | 手册类型 | ERC类型 | 结果

两侧来源互相独立：

    要求（手册名 / 手册类型）   pdf_pins.py 从数据手册 Table 5-1 抽取
    实测（网表名 / ERC类型）    kicad-cli sch export netlist 的引脚名
                                + kicad-cli sch erc 报告的引脚电气类型

**两边都不取自 .kicad_sym 的文本**，所以这张表能抓到"文本看着对、KiCad 读出来不对"
的错。（.kicad_sym 只用于搭建原理图，不参与取值。）

只用 openpyxl，不需要 Pillow —— 没有图片列。
"""

import os

__all__ = ["build_xlsx", "verdict", "HEAD"]

HEAD = ["pin", "手册名", "网表名", "手册类型", "ERC类型", "结果"]
COLW = [8, 16, 16, 18, 18, 10]
RES_FILL = {"PASS": ("C6EFCE", "006100"), "NG": ("FFC7CE", "9C0006")}
# openpyxl 的颜色不带 "#"


def verdict(row):
    """一行引脚比对的判定。

    只有"两边都有、名字一致、类型一致"才算 PASS。
    任一侧缺这个引脚 = NG（少引脚或多引脚，都是真错）。
    """
    has_pdf = row.get("pdf_name") not in (None, "", "-")
    has_net = row.get("netlist_name") not in (None, "", "-")
    if not has_pdf or not has_net:
        return "NG"
    if not row.get("name_ok") or not row.get("type_ok"):
        return "NG"
    return "PASS"


def build_xlsx(rows, out, title="", notes=None, sheet_name="引脚比对"):
    """rows: [{'pin','pdf_name','netlist_name','pdf_etype','erc_etype', ...}]"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEAD)
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[1].height = 24

    n_pass = n_ng = 0
    for i, row in enumerate(rows):
        r = i + 2
        res = verdict(row)
        n_pass += res == "PASS"
        n_ng += res == "NG"
        vals = [row.get("pin", i + 1),
                row.get("pdf_name") or "-",
                row.get("netlist_name") or "-",
                row.get("pdf_etype") or "-",
                row.get("erc_etype") or "-"]
        for k, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=k, value=v)
            cell.alignment = Alignment(
                horizontal="right" if k == 1 else "center", vertical="center")
            cell.border = border
        rc = ws.cell(row=r, column=6, value=res)
        bg, fg = RES_FILL.get(res, ("EDEDED", "808080"))
        rc.font = Font(bold=True, size=14, color=fg)
        rc.fill = PatternFill("solid", fgColor=bg)
        rc.alignment = Alignment(horizontal="center", vertical="center")
        rc.border = border
        ws.row_dimensions[r].height = 20

    for col, w in zip("ABCDEF", COLW):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    n = len(rows) + 3
    head = ("%s   ——   共 %d 个引脚，PASS %d，NG %d"
            % (title or "引脚比对", len(rows), n_pass, n_ng))
    c = ws.cell(row=n - 1, column=1, value=head)
    c.font = Font(bold=True, size=11,
                  color="006100" if not n_ng else "9C0006")
    for k, t in enumerate(notes or []):
        cc = ws.cell(row=n + k, column=1, value=t)
        cc.font = Font(size=9, italic=True, color="555555")
        cc.alignment = Alignment(vertical="center", wrap_text=True)

    try:
        wb.save(out)
    except PermissionError:
        base, ext = os.path.splitext(out)
        out = base + "(new)" + ext
        wb.save(out)
        print("note: 目标文件正被 Excel/WPS 打开，已另存为", out)
    return out


def main():
    import argparse
    import json

    ap = argparse.ArgumentParser(description="由 cmp_pins.json 生成引脚比对 xlsx")
    ap.add_argument("json_file")
    ap.add_argument("--out")
    ap.add_argument("--title", default="")
    a = ap.parse_args()
    rows = json.load(open(a.json_file, encoding="utf-8"))
    out = a.out or os.path.splitext(a.json_file)[0] + ".xlsx"
    p = build_xlsx(rows, out, a.title)
    print("report:", p)
    print("PASS %d / NG %d" % (sum(verdict(r) == "PASS" for r in rows),
                               sum(verdict(r) == "NG" for r in rows)))


if __name__ == "__main__":
    main()
